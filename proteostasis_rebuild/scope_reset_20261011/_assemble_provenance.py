from pathlib import Path
import csv, hashlib, re
from openpyxl import load_workbook
D=Path(__file__).resolve().parent
R=D.parent
T='experiment_grounded_rebuild_20261010/'
N=T+'novelty_audit/sources/'
rows=[]; paths=set()
def add(cid,claim,status,path,anchor,lines,lit,limit):
    p=(R/path).resolve(); text=p.read_text(); ls=text.splitlines()
    start=next(i for i,l in enumerate(ls,1) if anchor in l)
    end=min(len(ls),start+lines-1)
    paths.add(p)
    rows.append([cid,claim,status,str(p),start,end,lit,limit,ls[start-1][:240]])
add('C01','Original objective and canonical authors','observed','../envelope-paper/manuscript/MANUSCRIPT.md','Hannah E.',1,'Canonical manuscript','Authorship preserved; no claim of approval')
add('C01','Original two questions','observed','../envelope-paper/manuscript/MANUSCRIPT.md','First, is the envelope',1,'Canonical manuscript','Manuscript quotation, not historical user quotation')
add('C02','Error-attributable outcome subtracts error-free baseline','derived','scope_reset_20261011/MINIMAL_PHYSICAL_MODEL.md','Let `E=e`',17,'Causal total probability; LA24 supplies no outcome measurement','Common environment and fixed synthesis; not measured')
add('C03','Site-time times synthesis yields machine occupancy','derived','scope_reset_20261011/MINIMAL_PHYSICAL_MODEL.md','Separate four levels',10,'Little\'s law; FoldEco (FE12), FoldEcoSlim (FS19) contextual prior art','Finite stationary residence; repeat visits; background occupancy')
add('C04','Conservative finite-pool realization','derived','scope_reset_20261011/MINIMAL_PHYSICAL_MODEL.md','For clients `i`',26,'FoldEco (FE12); FoldEcoSlim (FS19), prior art not new proof','One site per chain, effective loss, no machinery decay beyond dilution')
add('C04','Older conservative framework and capped finite pools','derived',T+'theory/EQUATIONS.md','## A1.',103,'FoldEco (FE12)','Use reset equations for execution; do not import all prior pathways')
add('C05','Positive dilution bounds client mass; no generic runaway','derived','scope_reset_20261011/MINIMAL_PHYSICAL_MODEL.md','`N_req,i` must',1,'Conservation comparison inequality','Positive constant dilution; zero dilution needs removal')
add('C06','Steady occupancy identity is necessary','derived',T+'theory/EQUATIONS.md','## A4.',17,'Little\'s law; Schmidt (S2016)','No in-vivo J/tau in retrieved record; not global absence claim')
add('C07','Native output must replace dilution unfolding and degradation','derived',T+'theory/EQUATIONS.md','## A3.',17,'Santra et al. (SD17), context','Measured activity needed; conformation not necessarily function')
add('C08','Measured-input regulatory closure is conditional','conditional','scope_reset_20261011/MINIMAL_PHYSICAL_MODEL.md','If the observation window',1,'Tomoyasu (T98); Guisbert (G04); Lim (L13); Zhang (Z14)','Not calibrated or a prediction of sigma32')
add('C09','Feasible capacity does not ensure realized yield','derived',T+'theory/EQUATIONS.md','**Theorem 2',18,'FoldME (FM17), comparison only','Counterexample illustrative; published optimization costs not excluded')
add('C09','More capture need not improve native yield','derived',T+'kinetics/RESULTS.md','1. **Route criterion.',6,'Fauvet (H2021), motivating observation not validation','Fixed-pool analytic assumptions; no universal shared-pool sign')
add('C10','Conditional GroEL bypass exists','observed',T+'theory/EXPERIMENT_CONSTRAINTS.tsv','V04\t',1,'Vorderwulbecke 2004 (V04), DOI 10.1016/S0014-5793(04)00052-3','Duplicate-publication flag; used to construct model')
add('C10','DnaK association and aggregation patterns','observed',T+'theory/EXPERIMENT_CONSTRAINTS.tsv','C12a\t',3,'Calloni 2012 (C2012), DOI 10.1016/j.celrep.2011.12.007','Identity counts not flux; original client tables unavailable')
add('C11','Sigma32 feedback and TF exclusion','observed',T+'feedback/EXPERIMENTAL_CONTROL.md','| 1. Stress',9,'T98; G04; L13; Z14','Some sources abstract-level; no gain or functional delay estimate')
add('C11','Engineered regulation improves selected client yields','observed',T+'feedback/EXPERIMENTAL_CONTROL.md','## What the 2014 experiment',10,'Zhang 2014 (Z14), DOI 10.1021/cb5004477','Engineered BL21 state, not endogenous ceiling')
add('C11','Abundance is not throughput','observed','WO-07/REPORT.md','### EXPONENTIAL',12,'Schmidt 2016 (S2016), DOI 10.1038/nbt.3418','Matched abundance does not match kinetics; synthesis closure derived')
add('C12','Specific synonymous substitutions alter mDHFR GroEL dependence','observed',N+'D_htpg_codon_series/reingewertz2024.txt','Synonymous and non',1,'Reingewertz 2024 (R24), DOI 10.1002/pro.5087','Flattened primary text; timing computational, heterologous client')
add('C12','AGC contrast survives only a specified pipeline','observed','paper_reassessment_20261010/codon_bridge/CODON_RESULT.md','**Survives, narrowed:',7,'Landerer 2024 (LA24), PMC10939442; local corrected reanalysis','No portable amplitude or folding readout; localization artefact not excluded')
add('C13','Detection-conditional rate and broad synonymous comparisons withdrawn','withdrawn','paper_reassessment_20261010/codon_bridge/CODON_RESULT.md','**Withdraw as stated:',12,'Landerer 2024 (LA24), local reanalysis','Replacement standing frequencies do not identify nascent error rates')
add('C14','All-client recovery normalization did not test intended interaction','withdrawn','paper_reassessment_20261010/SCIENTIFIC_REPORT.md','The shipped route model fails',8,'V04; local route confrontation','ClpB requirement imposed; old simulation not biology')
add('C15','Fauvet point-estimate 20.5 to 18.9 percent','observed',T+'kinetics/fauvet2021_PMC8082187_fulltext.txt','18.5, 20.5, and 18.9%',1,'Fauvet 2021 (H2021), DOI 10.3389/fmolb.2021.653073','30 C insolubility not pure aggregate; total uncertainty unresolved')
add('C15','Kinetic model fails external point-estimate sign','conditional',T+'kinetics/RESULTS.md','**The decisive test',18,'Fauvet 2021 (H2021)','Archived 200/200 and 199/200 model samples; not rerun; old retrieval claim superseded')
add('C15','Later supplement provides per-protein means and SDs','observed',N+'D_htpg_codon_series/fauvet_supp_tables_headers.txt','Table_3.xlsx',3,'Fauvet 2021 (H2021), Table S3','No covariance or per-replicate total from summary SDs')
add('C16','Physiological headroom and legacy result classes rejected','withdrawn',T+'theory/CLAIM_DECISIONS.md','| usage-weighted',17,'LA24; FE12; H2021; earlier local audits','No replacement x63-113 or other physiological margin')
add('C16','Metal deployment and cross-species operational geometry rejected','withdrawn','../README.md','- **Metal-binding-site',6,'Historical root README','Historical claim disposition only; no active biological result')
add('C17','Capped reduction has finite states and only conditional hysteresis','derived',T+'theory/EQUATIONS.md','**Theorem A6',31,'FoldEco (FE12) contextual prior art','Different reduced model, parked; no physiological fold estimate')
add('C18','Finite chaperone network and growth models are prior art','observed',T+'theory/PRIOR_ART.md','| **FoldEco**',9,'Powers 2012 FE12; Santra 2017 SD17; Chen 2017 FM17','Historical review incomplete without FS19; no priority inferred from omissions')
add('C18','FoldEcoSlim already includes growth and dynamic stress response','observed',N+'C_foldecoslim/PMC6566562_main_text_pmc_html.txt','FoldEcoSlim also includes a dynamic response',5,'Pobre et al. 2019 (FS19), DOI 10.1002/pro.3639','Primary text; dynamic response and background are not new')
add('C19','Generic Hall subset counting fails for serial resource routes','derived',T+'constant_demand_v1/DERIVATION.md','We do not use a Hall',3,'Local mathematical correction; FM17 context','Whole-function demand branch and illustrative numbers excluded')
add('C20','No phase-matched kinetic bundle for physiological prediction','unsupported','WO-07/REPORT.md','**No. Neither bundle',25,'Schmidt 2016 S2016; local WO-07 source audit','Unsupported claim is calibrated physiological prediction, not bookkeeping')
add('C21','Competing error-workload predictions are prospective','conditional','scope_reset_20261011/DECISION_GATES.md','| **1.',4,'LA24, R24 and sigma32 studies motivate but do not test it','No experiment run; assay effect size and uncertainty must be prespecified')
add('C22','Interaction needs mechanistic attribution and Fauvet uncertainty','conditional','scope_reset_20261011/DECISION_GATES.md','| **4.',2,'Fauvet H2021; local gates','Nonzero I is not saturation; later supplement availability corrected')
add('C22','Babu held-out test is blocked separately','unsupported','paper_reassessment_20261010/SCIENTIFIC_REPORT.md','Serial and parallel route structure',7,'Babu 2014, DOI 10.1371/journal.pgen.1004120','Missing table and significant-hits-only censoring; no missing=neutral inference')
with (D/'CLAIM_AND_EVIDENCE_MATRIX.tsv').open('w') as f:
    w=csv.writer(f,delimiter='\t'); w.writerow(['claim_id','claim','status','source_file','line_start','line_end','named_literature_and_ids','limits','anchor_excerpt']); w.writerows(rows)
# Read binary supplements directly; inspect schema, do not invent aggregate uncertainty.
notes=['# Evidence reconciliation: later sources and exact remaining blockers','',
       '11 October 2026. This note supersedes blanket supplement-unavailable language in historical records. No fitted result or new biological measurement is introduced.','',
       '## Fauvet supplement inspection','',
       'The following locally archived original workbooks were opened read-only. Dimensions and header descriptions identify summary statistics, not replicate-level total insolubility.','']
for filename in ['Table_2.xlsx','Table_3.xlsx','Table_4.xlsx']:
    p=R/N/'D_htpg_codon_series/dl_fauvet'/filename; paths.add(p.resolve())
    wb=load_workbook(p,read_only=True,data_only=True)
    for sh in wb:
        notes.append(f'- `{filename}` / `{sh.title}`: {sh.max_row} rows × {sh.max_column} columns.')
        first=list(sh.iter_rows(min_row=1,max_row=min(2,sh.max_row),values_only=True))
        for row in first:
            nonempty=[str(v).replace('\n',' ') for v in row if v is not None]
            if nonempty: notes.append('  Header: '+ ' | '.join(nonempty))
    wb.close()
notes += ['', 'For a total T=ΣX_i, Var(T)=ΣVar(X_i)+2Σ_(i<j)Cov(X_i,X_j). Per-protein SDs do not identify the covariances. The between-strain contrast also needs the sampling/pairing design. No independence approximation, aggregate p-value, or synthetic replicate reconstruction is used here.', '',
          'The main text reports 20.5% versus 18.9% insolubility at 30 °C. Its proteomic normalization and membrane-protein component mean this is not identical to the model aggregate state. Supplementary Figure S1 describes growth at 37 °C; it does not supply a matched 30 °C rate. The missing evidence is total-fraction replicate uncertainty/covariance plus matched 30 °C growth and synthesis.', '',
          '## Scope and prior-art reconciliation','',
          'The original reset excluded manuscript drafting at that stage. The current user explicitly requested a full manuscript, so that earlier exclusion is superseded; the scientific boundary remains translation-associated burden and response.', '',
          'The earlier reassessment’s replacement ×63–113 margins remain rejected. Correcting an error-frequency input does not supply physiological capacity calibration.', '',
          'FoldEcoSlim (Pobre et al., DOI 10.1002/pro.3639, PMC6566562) already contains growth dilution, background proteome and sigma32-linked induction. The archived erratum PMC6863698 concerns acknowledgments. No novelty claim is based on adding those mechanisms.', '',
          'The theory constraint ledger contains stale labels referring to runaway or universal absence of in-vivo rates. This delivery uses only the bounded corrected equations and states that requisite rates are absent from the retrieved matched records; it does not claim nobody has ever measured them.', '',
          'All original files remain read-only historical evidence. Claims in old root README and model reports are not instructions to restore withdrawn results.']
(D/'EVIDENCE_RECONCILIATION.md').write_text('\n'.join(notes)+'\n')
# Preserve original verbatim, no header added to its bytes.
original=(R/'../envelope-paper/manuscript/MANUSCRIPT.md').resolve()
(D/'ORIGINAL_FROZEN_MANUSCRIPT.md').write_bytes(original.read_bytes())
# Locally reconcile historic scope notes, without changing original claim records elsewhere.
p=D/'ORIGIN_AND_SCOPE.md'
s=p.read_text().replace('and manuscript drafting. Synonymous codons', 'and new theory branches. Synonymous codons')
s += '\n## Completion update, 2026-10-11\n\nThe current user request explicitly authorized manuscript drafting, superseding the earlier stage-specific exclusion. The full deliverables are mapped in [README.md](README.md). The later Fauvet supplements are locally available as per-protein summary tables; the outstanding blocker is total-contrast uncertainty/covariance and matched 30 °C rates, as detailed in [EVIDENCE_RECONCILIATION.md](EVIDENCE_RECONCILIATION.md). Historical retrieval statements above describe the earlier audit stage. Reconstruction is finished; empirical validation remains unavailable.\n'
p.write_text(s)
p=D/'DECISION_GATES.md'; p.write_text(p.read_text()+'\n## Source-status clarification, 2026-10-11\n\nLater Fauvet Tables S2–S4 and the supplementary PDF are available in the novelty archive. They contain per-protein summary statistics; the missing gate-5 evidence is replicate-level total-fraction uncertainty/covariance and matched 30 °C growth/synthesis, not every supplement. See [EVIDENCE_RECONCILIATION.md](EVIDENCE_RECONCILIATION.md). These gates were specified, not run as biological experiments.\n')
# Include all user-requested context and primary prior-art manifests.
extra=['scope_reset_20261011/ORIGIN_AND_SCOPE.md','scope_reset_20261011/DECISION_GATES.md',
       'paper_reassessment_20261010/evidence/source_manifest.tsv',T+'feedback/CORRECTIONS.md',
       N+'C_foldecoslim/MANIFEST.txt',N+'C_foldecoslim/PMC6863698_erratum.txt',
       N+'D_htpg_codon_series/fauvet_Data_Sheet_1.txt']
for path in extra: paths.add((R/path).resolve())
with (D/'SOURCE_MANIFEST.tsv').open('w') as f:
    w=csv.writer(f,delimiter='\t');w.writerow(['source_file','bytes','sha256','text_lines'])
    for p in sorted(paths):
        b=p.read_bytes()
        try: n=len(b.decode().splitlines())
        except UnicodeDecodeError: n='binary'
        w.writerow([str(p),len(b),hashlib.sha256(b).hexdigest(),n])
(D/'LEDGER_PENDING.md').write_text('''# Append-ready ledger entry (not written externally)\n\nUser instruction limits all writes to this directory. The external append-only idea ledger was read but not modified. This entry records the synthesis decision and source-status correction; it is not a new wet-lab result.\n\n2026-10-11 | codex | DECISION | Proteostasis reconstruction completed as theoretical reassessment, not submission-ready; no physiological headroom; Fauvet supplements are present but total-contrast uncertainty and matched 30 C rates remain unresolved; FoldEcoSlim already includes background proteome and sigma32 response | proteostasis_law/proteostasis_rebuild/scope_reset_20261011 | Proteostasis envelope and fold condition\n''')
print('Matrix rows:',len(rows),'claim IDs:',len(set(r[0] for r in rows)))
print('Source files:',len(paths))
